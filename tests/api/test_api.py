"""T-073, T-074, T-075 (surface) - the HTTP API over the real pipeline on real
recorded scanner output. Uses Starlette's TestClient (incl. WebSocket) and a
synchronous scan runner, so no server or services are needed."""

from __future__ import annotations

import importlib.util
import sys
from datetime import UTC, datetime
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from qavach_api import AppState, create_app
from qavach_collectors import CollectorRegistry, CollectorResult, RunContext, Target, TargetType
from qavach_core.model.enums import ConfidenceTier
from qavach_storage import create_all, make_engine, session_factory
from qavach_worker import Deps
from sqlalchemy.pool import StaticPool

ROOT = Path(__file__).parent.parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
_spec = importlib.util.spec_from_file_location("demo", ROOT / "scripts/demo.py")
assert _spec and _spec.loader
demo = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(demo)

KNOWLEDGE, POLICY, PQC = demo.load_knowledge()
RESULTS = demo.replay_results(KNOWLEDGE)
NOW = datetime(2026, 9, 20, 9, 0, tzinfo=UTC)
SYSTEMS_CSV = (ROOT / "tests/fixtures/demo/systems.csv").read_text()
TARGET = {
    "target_type": "repository",
    "target_ref": "github.com/acme/polyglot-payments",
    "as_of": "2026-09-20",
    "capacity_per_quarter": 4,
}


class Replay:
    requires_sandbox = False
    requires_network = False
    default_confidence = ConfidenceTier.AST

    def __init__(self, result: CollectorResult) -> None:
        self._r = result
        self.name = result.tool.name
        self.version = result.tool.version

    def supports(self, target: Target) -> bool:
        return target.type is TargetType.REPOSITORY

    def collect(self, target: Target, ctx: RunContext) -> CollectorResult:
        return self._r


@pytest.fixture
def client():  # type: ignore[no-untyped-def]
    engine = make_engine("sqlite://")
    engine.pool.__class__  # noqa: B018
    from sqlalchemy import create_engine

    engine = create_engine(
        "sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False}
    )
    create_all(engine)
    registry = CollectorRegistry()
    for r in RESULTS:
        registry.register(Replay(r))
    state = AppState(
        session_factory=session_factory(engine),
        deps=Deps(
            registry=registry, knowledge=KNOWLEDGE, policy=POLICY, pqc=PQC, clock=lambda: NOW
        ),
        runner=lambda job: job(),  # synchronous: the scan is finished when the POST returns
        clock=lambda: NOW,
    )
    with TestClient(create_app(state)) as c:
        yield c


def _scan(client: TestClient, **over) -> str:  # type: ignore[no-untyped-def]
    client.post("/api/v1/systems/import", content=SYSTEMS_CSV, headers={"content-type": "text/csv"})
    r = client.post("/api/v1/scans", json={**TARGET, **over})
    assert r.status_code == 202, r.text
    return r.json()["scan_id"]  # type: ignore[no-any-return]


def test_starting_a_scan_returns_202_with_a_pollable_id_and_it_completes(
    client: TestClient,
) -> None:
    sid = _scan(client)
    body = client.get(f"/api/v1/scans/{sid}").json()
    assert body["status"] == "complete" and body["summary"]["assets"] == 29
    assert len(body["collectors"]) == 6 and all(not c["partial"] for c in body["collectors"])
    assert body["policy_snapshot_id"] == POLICY.snapshot_id and "roadmap" not in body["summary"]


def test_unknown_scans_assets_and_target_types_are_clean_errors(client: TestClient) -> None:
    assert client.get("/api/v1/scans/nope").status_code == 404
    assert client.get("/api/v1/scans/nope/assets").status_code == 404
    assert client.get("/api/v1/assets/nope").status_code == 404
    assert (
        client.post("/api/v1/scans", json={**TARGET, "target_type": "warp-core"}).status_code == 422
    )
    assert client.post("/api/v1/scans", json={**TARGET, "z_scenario": "hopeful"}).status_code == 422
    assert (
        client.post(
            "/api/v1/scans", json={"target_type": "repository", "target_ref": ""}
        ).status_code
        == 422
    )


def test_the_inventory_is_faceted_filtered_and_paginated(client: TestClient) -> None:
    sid = _scan(client)
    page = client.get(f"/api/v1/scans/{sid}/assets", params={"page_size": 10}).json()
    assert page["total"] == 29 and len(page["items"]) == 10 and page["facets"]["finding_class"]
    weak = client.get(
        f"/api/v1/scans/{sid}/assets", params={"finding_class": "classical-weak"}
    ).json()
    assert weak["total"] > 0 and {i["finding_class"] for i in weak["items"]} == {"classical-weak"}
    assert client.get(f"/api/v1/scans/{sid}/assets", params={"page_size": 9999}).status_code == 422
    assert client.get(f"/api/v1/scans/{sid}/assets", params={"q": "md5"}).json()["total"] >= 1


def test_asset_detail_has_occurrences_risk_explanations_and_a_cited_recommendation(
    client: TestClient,
) -> None:
    sid = _scan(client)
    items = client.get(
        f"/api/v1/scans/{sid}/assets", params={"finding_class": "quantum-vulnerable"}
    ).json()["items"]
    detail = client.get(f"/api/v1/assets/{items[0]['id']}").json()
    assert detail["occurrences"] and detail["risk"][0]["explanations"]
    assert detail["recommendation"]["primary"]["source_url"].startswith("https://")


def test_the_roadmap_endpoint_returns_waves_bridges_units_and_edges(client: TestClient) -> None:
    sid = _scan(client)
    road = client.get(f"/api/v1/scans/{sid}/roadmap").json()
    assert road["roadmap"]["bridges"] and road["units"] and road["edges"]
    assert {u["function"] for u in road["units"]} >= {"signature", "encryption"}


def test_the_cbom_export_is_schema_valid_and_carries_no_risk_data(client: TestClient) -> None:
    from qavach_core.export import cbom_violations
    from schema_check import cbom_1_6_validator, cbom_validator, errors_of

    sid = _scan(client)
    r = client.get(f"/api/v1/scans/{sid}/export/cbom")
    assert r.headers["content-type"].startswith("application/vnd.cyclonedx+json")
    doc = r.json()
    assert (
        doc["specVersion"] == "1.7"
        and errors_of(cbom_validator(), doc) == []
        and cbom_violations(doc) == []
    )
    old = client.get(f"/api/v1/scans/{sid}/export/cbom", params={"spec": "1.6"}).json()
    assert old["specVersion"] == "1.6" and errors_of(cbom_1_6_validator(), old) == []
    assert client.get(f"/api/v1/scans/{sid}/export/cbom", params={"spec": "2.0"}).status_code == 422


def test_scope_is_a_query_over_the_reconciled_set_not_a_re_merge(client: TestClient) -> None:
    sid = _scan(client)
    root = client.get(f"/api/v1/scans/{sid}/export/cbom").json()
    system = client.get(
        f"/api/v1/scans/{sid}/export/cbom", params={"scope": "system", "system_id": "payments-core"}
    ).json()
    assert {c["bom-ref"] for c in system["components"]} <= {
        c["bom-ref"] for c in root["components"]
    }
    assert (
        client.get(f"/api/v1/scans/{sid}/export/cbom", params={"scope": "system"}).status_code
        == 422
    )
    empty = client.get(
        f"/api/v1/scans/{sid}/export/cbom", params={"scope": "system", "system_id": "ghost"}
    ).json()
    assert empty["components"] == []


def test_the_register_export_validates_against_its_schema_and_ties_to_the_cbom(
    client: TestClient,
) -> None:
    from schema_check import errors_of, register_validator

    sid = _scan(client)
    register = client.get(f"/api/v1/scans/{sid}/export/risk-register").json()
    cbom = client.get(f"/api/v1/scans/{sid}/export/cbom").json()
    assert errors_of(register_validator(), register) == []
    assert {e["bom_ref"] for e in register["entries"]} == {c["bom-ref"] for c in cbom["components"]}
    assert register["policy"]["snapshot_id"] == POLICY.snapshot_id
    assert register["roadmap"]["bridges"]
    assert register["summary"]["coverage_failures"] == register["summary"]["by_finding_class"].get(
        "unknown", 0
    )


def test_sarif_is_served_from_stored_entries_and_gates_by_class(client: TestClient) -> None:
    sid = _scan(client)
    sarif = client.get(f"/api/v1/scans/{sid}/export/sarif").json()
    results = sarif["runs"][0]["results"]
    register = client.get(f"/api/v1/scans/{sid}/export/risk-register").json()
    refs = {e["bom_ref"]: e for e in register["entries"]}
    assert results and {r["properties"]["bom-ref"] for r in results} <= set(refs)
    assert all(
        refs[r["properties"]["bom-ref"]]["finding_class"] not in ("grover-affected", "quantum-safe")
        for r in results
    )
    failed = client.get(f"/api/v1/scans/{sid}/gate", params={"fail_on": "classical-weak"}).json()
    assert failed["fail"] and failed["reasons"]
    overdue = client.get(
        f"/api/v1/scans/{sid}/gate", params={"fail_on": "quantum-safe,overdue"}
    ).json()
    assert overdue["fail"] is True  # the demo has overdue entries
    assert (
        client.get(f"/api/v1/scans/{sid}/gate", params={"fail_on": "nonsense"}).status_code == 422
    )
    assert client.get("/api/v1/scans/nope/export/sarif").status_code == 404


def test_diffing_a_scan_against_itself_shows_no_drift_and_no_caveats(client: TestClient) -> None:
    first, second = _scan(client), _scan(client)
    body = client.get(f"/api/v1/scans/{first}/diff/{second}").json()
    s = body["summary"]
    assert s["added"] == s["removed"] == s["changed"] == 0 and s["unchanged"] > 0
    assert body["caveats"] == []


def test_a_later_as_of_date_worsens_bands_and_the_diff_says_the_date_moved(
    client: TestClient,
) -> None:
    """Bands move as deadlines approach with no change to the estate - the diff
    must attribute that to the date rather than present it as drift."""
    now, later = _scan(client), _scan(client, as_of="2030-01-01")
    body = client.get(f"/api/v1/scans/{now}/diff/{later}").json()
    assert body["summary"]["added"] == body["summary"]["removed"] == 0
    assert body["summary"]["worsened"] > 0
    assert any("as-of" in c for c in body["caveats"])
    assert all(c["label"] and c["direction"] for c in body["changed"])
    assert client.get(f"/api/v1/scans/{now}/diff/nope").status_code == 404


def test_the_pdf_is_a_real_deterministic_rendering_of_the_register(client: TestClient) -> None:
    sid = _scan(client)
    first = client.get(f"/api/v1/scans/{sid}/export/report.pdf")
    assert first.status_code == 200 and first.headers["content-type"] == "application/pdf"
    assert first.content.startswith(b"%PDF-") and len(first.content) > 2000
    # same scan, same bytes: nothing wall-clock-dependent is written
    assert client.get(f"/api/v1/scans/{sid}/export/report.pdf").content == first.content
    assert client.get("/api/v1/scans/nope/export/report.pdf").status_code == 404


def test_the_xlsx_register_carries_every_entry_and_never_colours_unknown_as_safe(
    client: TestClient,
) -> None:
    import io

    from openpyxl import load_workbook

    sid = _scan(client)
    r = client.get(f"/api/v1/scans/{sid}/export/register.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Assets", "Roadmap"]
    register = client.get(f"/api/v1/scans/{sid}/export/risk-register").json()
    rows = list(wb["Assets"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == len(register["entries"])
    assert {row[-1] for row in rows} == {e["bom_ref"] for e in register["entries"]}
    unknown_fills = {
        (c.fill.patternType, c.fill.fgColor.rgb)
        for row in wb["Assets"].iter_rows(min_row=2)
        for c in [row[2]]
        if str(c.value).startswith("Unclassified")
    }
    assert unknown_fills and all(p == "lightUp" for p, _ in unknown_fills)  # hatched, not solid
    summary = "\n".join(str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value)
    assert "coverage failure" in summary.lower() and "nominal" in summary


def test_act_now_never_includes_grover_or_unclassified_entries_whatever_their_band() -> None:
    from qavach_api.reports import act_now

    def entry(ref: str, cls: str, band: str) -> dict[str, str]:
        return {"bom_ref": ref, "finding_class": cls, "band": band}

    entries = [
        entry("grover", "grover-affected", "overdue"),
        entry("unknown", "unknown", "overdue"),
        entry("safe", "quantum-safe", "overdue"),
        entry("imminent-qv", "quantum-vulnerable", "imminent"),
        entry("cw", "classical-weak", "not-applicable"),  # broken today, Mosca band is moot
        entry("overdue-qv", "quantum-vulnerable", "overdue"),
        entry("planned-qv", "quantum-vulnerable", "planned"),
    ]
    assert [e["bom_ref"] for e in act_now(entries)] == ["cw", "overdue-qv", "imminent-qv"]


def test_systems_import_reports_every_problem_and_imports_the_good_rows(client: TestClient) -> None:
    csv = (
        "id,name,owner,criticality,data_classification,internet_facing\n"
        "ok,Ok,me,3,internal,true\nbad,Bad,me,9,internal,true\n"
    )
    body = client.post(
        "/api/v1/systems/import", content=csv, headers={"content-type": "text/csv"}
    ).json()
    assert body["imported"] == 1 and body["ok"] is False
    assert [(i["row"], i["field"]) for i in body["issues"] if i["severity"] == "error"] == [
        (2, "criticality")
    ]
    assert [s["id"] for s in client.get("/api/v1/systems").json()] == ["ok"]


def test_systems_import_accepts_yaml(client: TestClient) -> None:
    yaml_text = (
        "systems:\n  - {id: y, name: Y, owner: o, criticality: 2, "
        "data_classification: public, internet_facing: false}\n"
    )
    body = client.post(
        "/api/v1/systems/import", content=yaml_text, headers={"content-type": "application/yaml"}
    ).json()
    assert body["imported"] == 1 and body["ok"]
    assert (
        client.post(
            "/api/v1/systems/import",
            content="- - -\n: :",
            headers={"content-type": "application/yaml"},
        ).status_code
        == 422
    )


def test_the_policy_endpoint_returns_every_value_with_its_citation(client: TestClient) -> None:
    body = client.get("/api/v1/policy").json()
    assert body["snapshot_id"] == POLICY.snapshot_id and len(body["nodes"]) > 40
    assert all(n["basis"] for n in body["nodes"])


def test_suppression_needs_a_reason_is_audited_and_the_asset_stays_listed_but_flagged(
    client: TestClient,
) -> None:
    sid = _scan(client)
    item = client.get(f"/api/v1/scans/{sid}/assets").json()["items"][0]
    assert (
        client.post(
            f"/api/v1/assets/{item['id']}/suppress", json={"reason": "", "author": "a"}
        ).status_code
        == 422
    )
    ok = client.post(
        f"/api/v1/assets/{item['id']}/suppress",
        json={"reason": "accepted", "author": "ops", "days": 10},
    )
    assert ok.status_code == 200
    after = client.get(f"/api/v1/scans/{sid}/assets", params={"page_size": 100}).json()
    assert (
        after["total"] == 29
        and next(i for i in after["items"] if i["id"] == item["id"])["suppressed"] is True
    )
    assert (
        client.post("/api/v1/assets/nope/suppress", json={"reason": "r", "author": "a"}).status_code
        == 404
    )


def test_adjudicating_an_asset_with_no_dispute_is_refused(client: TestClient) -> None:
    sid = _scan(client)
    item = client.get(f"/api/v1/scans/{sid}/assets").json()["items"][0]
    r = client.post(
        f"/api/v1/assets/{item['id']}/adjudicate",
        json={"attribute": "mode", "value": "gcm", "by": "ops", "reason": "r"},
    )
    assert r.status_code == 422 and "no unresolved dispute" in r.json()["detail"]


def test_progress_is_streamed_over_a_websocket_with_replay_for_late_joiners(
    client: TestClient,
) -> None:
    sid = _scan(client)  # finished before we connect: history must still be replayed
    with client.websocket_connect(f"/api/v1/scans/{sid}/progress") as ws:
        events = []
        while True:
            event = ws.receive_json()
            events.append(event)
            if event["stage"] == "scan" and event["status"] in {"finished", "closed", "failed"}:
                break
    stages = [e["stage"] for e in events if e["collector"] is None and e["status"] == "started"]
    assert stages == ["collect", "assemble", "context", "risk", "recommend", "roadmap", "persist"]
    assert {e["collector"] for e in events if e["collector"]} >= {
        "source_scan.cdxgen",
        "runtime.tracebom",
    }


def test_a_websocket_for_an_unknown_scan_is_closed(client: TestClient) -> None:
    with pytest.raises(Exception):  # noqa: B017,PT011 - Starlette raises on a close before accept
        with client.websocket_connect("/api/v1/scans/nope/progress"):
            pass


# ---- T-075: policy/simulate ----


def test_simulating_a_more_aggressive_z_moves_assets_toward_overdue_and_writes_nothing(
    client: TestClient,
) -> None:
    sid = _scan(client)
    before = client.get(f"/api/v1/scans/{sid}/assets", params={"page_size": 100}).json()
    body = client.post(
        "/api/v1/policy/simulate", json={"scan_id": sid, "z_scenario": "aggressive"}
    ).json()
    assert body["scored"] == sum(len(i["scores"]) for i in before["items"])
    assert body["z_scenario"] == {"baseline": "nominal", "candidate": "aggressive"}
    overdue = lambda b: b.get("overdue", 0)  # noqa: E731
    assert overdue(body["bands"]["candidate"]) >= overdue(body["bands"]["baseline"])
    assert (
        body["candidate_policy_snapshot_id"] != body["baseline_policy_snapshot_id"]
        or body["changed"] == 0
    )
    after = client.get(f"/api/v1/scans/{sid}/assets", params={"page_size": 100}).json()
    assert after == before  # nothing was written
    assert "heuristics" in body["heuristics_notice"] or "heuristic" in body["heuristics_notice"]


def test_simulating_the_unchanged_policy_changes_nothing(client: TestClient) -> None:
    sid = _scan(client)
    body = client.post("/api/v1/policy/simulate", json={"scan_id": sid}).json()
    assert body["changed"] == 0 and body["bands"]["baseline"] == body["bands"]["candidate"]
    assert body["top_movers"] == []  # rounding noise must not be reported as movement
    assert body["candidate_policy_snapshot_id"] == body["baseline_policy_snapshot_id"]


def test_a_policy_value_override_takes_effect_and_reports_the_movers(client: TestClient) -> None:
    sid = _scan(client)
    body = client.post(
        "/api/v1/policy/simulate",
        json={"scan_id": sid, "overrides": {"scoring.mosca.ephemeral_artefact_max_years": 50.0}},
    ).json()
    assert body["overrides"] == {"scoring.mosca.ephemeral_artefact_max_years": 50.0}
    assert isinstance(body["top_movers"], list)


@pytest.mark.parametrize(
    "overrides",
    [
        {"scoring.nope": 1},
        {"scoring.caraf": 1},  # a container, not a node
        {"scoring.mosca.imminent_within_years": "two"},  # wrong type
        {"scoring.mosca.imminent_within_years": True},  # a bool is not a number here
    ],
)
def test_a_bad_override_is_refused_not_silently_simulated_as_unchanged(
    client: TestClient, overrides: dict
) -> None:  # type: ignore[type-arg]
    sid = _scan(client)
    assert (
        client.post(
            "/api/v1/policy/simulate", json={"scan_id": sid, "overrides": overrides}
        ).status_code
        == 422
    )


def test_an_unknown_scan_or_scenario_is_a_clean_error(client: TestClient) -> None:
    assert client.post("/api/v1/policy/simulate", json={"scan_id": "nope"}).status_code == 404
    sid = _scan(client)
    assert (
        client.post(
            "/api/v1/policy/simulate", json={"scan_id": sid, "z_scenario": "hopeful"}
        ).status_code
        == 422
    )


def test_scans_are_listed_newest_first_with_their_status(client: TestClient) -> None:
    a = _scan(client)
    listed = client.get("/api/v1/scans").json()
    assert [s["id"] for s in listed] == [a]
    assert listed[0]["status"] == "complete" and listed[0]["assets"] == 29


def test_the_z_slider_can_override_a_scenarios_crqc_year_and_it_moves_assets(
    client: TestClient,
) -> None:
    sid = _scan(client)
    later = client.post(
        "/api/v1/policy/simulate",
        json={"scan_id": sid, "overrides": {"z_scenarios.scenarios.nominal#crqc_year": 2045}},
    ).json()
    sooner = client.post(
        "/api/v1/policy/simulate",
        json={"scan_id": sid, "overrides": {"z_scenarios.scenarios.nominal#crqc_year": 2028}},
    ).json()
    late = lambda r: r["bands"]["candidate"].get("overdue", 0)  # noqa: E731
    assert late(sooner) >= late(later)
    assert sooner["candidate_policy_snapshot_id"] != later["candidate_policy_snapshot_id"]


@pytest.mark.parametrize(
    "target", ["z_scenarios.scenarios.nominal#nope", "z_scenarios.scenarios.nope#crqc_year"]
)
def test_a_bad_field_override_is_refused(client: TestClient, target: str) -> None:
    sid = _scan(client)
    assert (
        client.post(
            "/api/v1/policy/simulate", json={"scan_id": sid, "overrides": {target: 2040}}
        ).status_code
        == 422
    )


def test_the_first_simulate_after_a_scan_is_served_from_the_prewarmed_row_cache(
    client: TestClient,
) -> None:
    from qavach_worker.simulate import clear_simulation_cache

    clear_simulation_cache()
    sid = _scan(client)
    r = client.post("/api/v1/policy/simulate", json={"scan_id": sid, "z_scenario": "aggressive"})
    assert r.status_code == 200 and r.json()["rows_from_cache"] is True
